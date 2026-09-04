"""Generate advisory person-image review records for D11.5B.

This tool reads the candidate workbook and the existing people text-review
records without editing either source.  It never downloads images, writes to
``frontend/public/images/people``, touches MySQL, or changes reviewed data.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


PEOPLE_SHEET = "P0-历史人物"
EXPECTED_PEOPLE_COUNT = 31
AUDIT_RETRIEVED_AT = "2026-09-04"

UNCERTAIN_IMAGE_TYPES = {
    "朱德": "portrait",
    "宣侠父": "content",
    "白求恩": "group",
    "柯棣华": "group",
    "埃德加·斯诺": "portrait",
    "海伦·福斯特·斯诺": "group",
    "史沫特莱": "portrait",
}

PUBLISHERS = {
    "朱德": "共产党员网 / 中共一大会址纪念馆藏（Excel 图片线索，未核验对应图片页）",
    "宣侠父": "八路军西安办事处纪念馆官网（Excel 图片线索，未核验身份说明）",
    "白求恩": "人民网党史频道（Excel 图片线索，未核验合影身份位置）",
    "柯棣华": "人民网党史频道（Excel 图片线索，未核验合影身份位置）",
    "埃德加·斯诺": "新华书店总店 / 新华出版社（Excel 图片线索，未核验对应人像页）",
    "海伦·福斯特·斯诺": "八路军西安办事处纪念馆官网（Excel 图片线索，未核验合影身份位置）",
    "史沫特莱": "人民网党史频道（Excel 图片线索，未核验对应肖像页）",
}

PINYIN = {
    "甲": "jia", "周": "zhou", "恩": "en", "来": "lai", "朱": "zhu", "德": "de", "刘": "liu", "少": "shao", "奇": "qi",
    "叶": "ye", "剑": "jian", "英": "ying", "林": "lin", "伯": "bo", "渠": "qu", "董": "dong", "必": "bi", "武": "wu",
    "伍": "wu", "云": "yun", "甫": "fu", "子": "zi", "健": "jian", "宣": "xuan", "侠": "xia", "父": "fu", "李": "li", "克": "ke", "农": "nong",
    "鼎": "ding", "冯": "feng", "海": "hai", "涂": "tu", "作": "zuo", "潮": "chao", "白": "bai", "求": "qiu", "柯": "ke", "棣": "di", "华": "hua",
    "埃": "edgar", "加": "jia", "斯": "s", "诺": "snow", "伦": "helen", "福": "foster", "特": "te", "史": "smedley", "沫": "mo", "莱": "lai",
    "陈": "chen", "嘉": "jia", "庚": "geng", "邓": "deng", "颖": "ying", "超": "chao", "康": "kang", "清": "qing", "博": "bo", "古": "gu", "秦": "qin", "邦": "bang", "宪": "xian",
    "彭": "peng", "怀": "huai", "小": "xiao", "平": "ping", "赓": "geng", "冼": "xian", "星": "xing", "王": "wang", "浩": "hao", "礼": "li", "高": "gao", "万": "wan", "季": "ji", "壮": "zhuang",
}


def read_people_rows(workbook_path: Path) -> list[dict]:
    """Return non-empty P0-历史人物 rows while leaving the workbook untouched."""
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook[PEOPLE_SHEET]
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if any(value is not None for value in row.values()):
            rows.append(row)
    if len(rows) != EXPECTED_PEOPLE_COUNT:
        raise ValueError(f"{PEOPLE_SHEET} expected {EXPECTED_PEOPLE_COUNT} records, got {len(rows)}")
    return rows


def image_source_level(person_review: dict) -> str:
    return person_review.get("sourceLevel") or "C"


def source_page_title(source_url: str) -> str | None:
    if "dangshi.people.com.cn" in source_url:
        return "七贤庄——抗日战争的红色驿站——走进八路军西安办事处纪念馆"
    if "xabb.org.cn" in source_url:
        return "八路军西安办事处纪念馆-展馆介绍"
    return None


def image_payload(
    status: str,
    *,
    source_page_url: str | None = None,
    publisher: str | None = None,
    page_title: str | None = None,
    identity_evidence: str | None = None,
    source_level: str | None = None,
    notes: str,
) -> dict:
    return {
        "status": status,
        "sourceUrl": None,
        "sourcePageUrl": source_page_url,
        "publisher": publisher,
        "pageTitle": page_title,
        "caption": None,
        "identityEvidence": identity_evidence,
        "sourceLevel": source_level,
        "retrievedAt": AUDIT_RETRIEVED_AT if source_page_url else None,
        "recommendedForUse": False,
        "notes": notes,
    }


def classify_image_lead(name: str, description: str, source_url: str, source_type: str, source_level: str = "C") -> dict:
    """Classify an Excel image lead without promoting it to a usable image."""
    description = (description or "").strip()
    source_url = (source_url or "").strip()
    common = {
        "source_page_url": source_url or None,
        "source_level": source_level,
        "page_title": source_page_title(source_url),
    }
    source_page_availability = (
        "unavailableDuringAudit" if "xxbcm.com" in source_url
        else "checkedNoIdentityEvidence" if source_url
        else "notApplicable"
    )

    if "外景" in description or "建筑" in description or "旧址" in description:
        return {
            "status": "missing",
            "image": image_payload(
                "missing",
                publisher=source_type or None,
                notes="现有线索为纪念馆/旧址外景，不是人物头像，不能绑定到人物。",
                **common,
            ),
            "relatedVisualAvailable": True,
            "relatedImageType": "site",
            "imageType": None,
            "suggestedUse": "notRecommended",
            "sourcePageAvailability": source_page_availability,
        }
    if "手迹" in description or "题字" in description or "题诗" in description:
        return {
            "status": "missing",
            "image": image_payload(
                "missing",
                publisher=source_type or None,
                notes="现有线索为题字或手迹，不是人物肖像；可在后续历史内容中另行审核为相关文物图片。",
                **common,
            ),
            "relatedVisualAvailable": True,
            "relatedImageType": "handwriting",
            "imageType": None,
            "suggestedUse": "relatedArtifact",
            "sourcePageAvailability": source_page_availability,
        }
    if name not in UNCERTAIN_IMAGE_TYPES and any(token in description for token in ("文物图片", "手表", "相机", "长袍")):
        return {
            "status": "missing",
            "image": image_payload(
                "missing",
                publisher=source_type or None,
                notes="现有线索是人物相关文物或物件，不是人物肖像，不能作为详情页头像。",
                **common,
            ),
            "relatedVisualAvailable": True,
            "relatedImageType": "artifact",
            "imageType": None,
            "suggestedUse": "relatedArtifact",
            "sourcePageAvailability": source_page_availability,
        }
    if name in UNCERTAIN_IMAGE_TYPES:
        image_type = UNCERTAIN_IMAGE_TYPES[name]
        if image_type == "group":
            identity = "Excel 仅提供合影线索，现有来源页未明确标出该候选人物在照片中的身份或位置。"
            notes = "合影不可作为头像；若未来来源页明确合影人物，可另作历史内容插图复审。"
            suggested_use = "contentImage"
        elif name == "宣侠父":
            identity = "Excel 提到人物书写照片，但现有来源页没有可核验的图片说明与人物一一对应关系。"
            notes = "可能相关的活动照片暂不作为头像；需补充可靠图片页和说明。"
            suggested_use = "contentImage"
        else:
            identity = "Excel 提供了可能的人像线索，但当前可追溯页面未提供可核验的图片说明与人物一一对应关系。"
            notes = "不得按文件名、搜索结果或常识确认身份；需补充可靠图片页及明确说明。"
            suggested_use = "notRecommended"
        if source_page_availability == "unavailableDuringAudit":
            notes = "来源页在本轮定向核验中因 TLS 证书验证失败而无法读取；未获得人物身份依据，不能作为头像来源。"
        return {
            "status": "uncertain",
            "image": image_payload(
                "uncertain",
                publisher=PUBLISHERS.get(name, source_type or None),
                identity_evidence=identity,
                notes=notes,
                **common,
            ),
            "relatedVisualAvailable": False,
            "relatedImageType": None,
            "imageType": image_type,
            "suggestedUse": suggested_use,
            "sourcePageAvailability": source_page_availability,
        }
    return {
        "status": "missing",
        "image": image_payload(
            "missing",
            notes="Excel 未提供可核验的人物图片线索；保持中性剪影，不以文字资料来源替代图片来源。",
        ),
        "relatedVisualAvailable": False,
        "relatedImageType": None,
        "imageType": None,
        "suggestedUse": "notRecommended",
        "sourcePageAvailability": source_page_availability,
    }


def slugify_name(name: str) -> str:
    parts = []
    for character in name:
        if character in PINYIN:
            parts.append(PINYIN[character])
        elif character.isascii() and character.isalnum():
            parts.append(character.lower())
    return "-".join(part for part in parts if part) or "person"


def localization_candidates(records: list[dict]) -> list[dict]:
    candidates = []
    for record in records:
        image = record["image"]
        if image["status"] == "confirmed" and image["recommendedForUse"]:
            candidates.append({
                "candidateId": record["candidateId"],
                "name": record["name"],
                "imageSourceUrl": image["sourceUrl"],
                "sourcePageUrl": image["sourcePageUrl"],
                "publisher": image["publisher"],
                "suggestedFilename": f"{slugify_name(record['name'])}.jpg",
                "approvedForDownload": False,
            })
    return candidates


def find_duplicate_image_bindings(records: list[dict]) -> list[dict]:
    groups: defaultdict[str, list[str]] = defaultdict(list)
    for record in records:
        image = record["image"]
        if image["status"] == "confirmed" and image["recommendedForUse"] and image["sourceUrl"]:
            groups[image["sourceUrl"]].append(record["candidateId"])
    return [
        {"sourceUrl": source_url, "candidateIds": candidate_ids}
        for source_url, candidate_ids in groups.items()
        if len(candidate_ids) > 1
    ]


def build_records(rows: list[dict], person_reviews: list[dict]) -> list[dict]:
    reviews_by_id = {record["candidateId"]: record for record in person_reviews}
    records = []
    for row in rows:
        candidate_id = f"PEOPLE-{int(row['序号']):03d}"
        person_review = reviews_by_id.get(candidate_id)
        if person_review is None:
            raise ValueError(f"Missing text-review record for {candidate_id}")
        name = str(row["姓名"] or "").strip()
        source_url = str(row.get("来源(URL)") or "").strip()
        source_type = str(row.get("来源类型") or "").strip()
        description = str(row.get("图片来源") or "").strip()
        classification = classify_image_lead(name, description, source_url, source_type, image_source_level(person_review))
        records.append({
            "candidateId": candidate_id,
            "name": name,
            "personRecommendedDecision": person_review["recommendedDecision"],
            "personReviewStatus": person_review["reviewStatus"],
            "originalImageSourceDescription": description,
            "originalSourceUrl": source_url or None,
            "originalSourceType": source_type or None,
            "originalVerificationStatus": row.get("核验状态"),
            "image": classification["image"],
            "imageType": classification["imageType"],
            "relatedVisualAvailable": classification["relatedVisualAvailable"],
            "relatedImageType": classification["relatedImageType"],
            "suggestedUse": classification["suggestedUse"],
            "sourcePageAvailability": classification["sourcePageAvailability"],
        })
    return records


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def update_qa_summary(qa_path: Path, records: list[dict]) -> dict:
    qa = json.loads(qa_path.read_text(encoding="utf-8")) if qa_path.exists() else {}
    statuses = Counter(record["image"]["status"] for record in records)
    uncertain = [record for record in records if record["image"]["status"] == "uncertain"]
    qa.update({
        "peopleImageConfirmed": statuses["confirmed"],
        "peopleImageUncertain": statuses["uncertain"],
        "peopleImageMissing": statuses["missing"],
        "peopleImageRejected": statuses["rejected"],
        "searchEngineOnlyImages": sum(
            1 for record in records
            if any(token in (record["image"]["sourceUrl"] or "").lower() for token in ("baidu", "bing", "google"))
        ),
        "identityUnclearImages": len(uncertain),
        "duplicateImageBindings": find_duplicate_image_bindings(records),
    })
    write_json(qa_path, qa)
    return qa


def write_markdown(path: Path, records: list[dict]) -> None:
    statuses = Counter(record["image"]["status"] for record in records)
    lines = [
        "# D11.5B 历史人物图片审核", "",
        "本清单只审核人物图片候选，不改变 `people-review.json` 中的人物文字审核建议。所有候选均未获下载授权；未下载图片、未修改 Excel、MySQL、游客端或正式 `reviewed/` 数据。", "",
        "## 审核原则", "",
        "- 只有可靠来源页同时明确图片主体与人物身份，才可标记 `confirmed` 并建议作为头像。",
        "- 建筑、旧址、手迹、题字与人物相关文物不是人物头像。合影若未明确人物身份或位置，最多保留为不推荐的内容插图线索。",
        "- 搜索引擎仅可用于发现原页面，不能作为图片来源；本批记录未登记搜索引擎缩略图。", "",
        "## 人物图片审核清单", "",
        "| ID | 人物 | 文字审核建议 | 人物照片状态 | 来源机构 | 身份明确 | 建议用途 |", "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for record in records:
        image = record["image"]
        publisher = (image["publisher"] or "-").replace("|", "／")
        identity = "是" if image["status"] == "confirmed" else "否"
        lines.append(
            f"| {record['candidateId']} | {record['name']} | {record['personRecommendedDecision']} | "
            f"{image['status']} | {publisher} | {identity} | {record['suggestedUse']} |"
        )
    lines.extend(["", "## 统计", "", f"- confirmed：{statuses['confirmed']}", f"- uncertain：{statuses['uncertain']}", f"- missing：{statuses['missing']}", f"- rejected：{statuses['rejected']}", ""])
    for status, label in (("confirmed", "已确认"), ("uncertain", "待确认"), ("missing", "缺失"), ("rejected", "排除")):
        names = "、".join(record["name"] for record in records if record["image"]["status"] == status) or "无"
        lines.append(f"- {label}人物：{names}")
    lines.extend([
        "", "## 后续边界", "",
        "本轮 `people-image-localization-candidates.json` 仅能接收 `confirmed` 且 `recommendedForUse: true` 的候选。即使进入清单，`approvedForDownload` 仍必须为 `false`，等待人工逐条批准后才可进入 D11.5C 本地化。",
    ])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def generate(workbook_path: Path, people_review_path: Path, output_dir: Path, docs_dir: Path) -> dict:
    rows = read_people_rows(workbook_path)
    person_reviews = json.loads(people_review_path.read_text(encoding="utf-8"))
    if len(person_reviews) != EXPECTED_PEOPLE_COUNT:
        raise ValueError(f"people-review expected {EXPECTED_PEOPLE_COUNT} records, got {len(person_reviews)}")
    output_dir.mkdir(parents=True, exist_ok=True)
    docs_dir.mkdir(parents=True, exist_ok=True)
    records = build_records(rows, person_reviews)
    localization = localization_candidates(records)
    write_json(output_dir / "people-image-review.json", records)
    write_json(output_dir / "people-image-localization-candidates.json", localization)
    qa = update_qa_summary(output_dir / "qa-summary.json", records)
    write_markdown(docs_dir / "people-image-review.md", records)
    return {"records": records, "localization": localization, "qa": qa}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--people-review", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--docs-dir", required=True, type=Path)
    args = parser.parse_args()
    result = generate(args.input, args.people_review, args.output_dir, args.docs_dir)
    print(json.dumps({
        "people": len(result["records"]),
        "localizationCandidates": len(result["localization"]),
        "peopleImageConfirmed": result["qa"]["peopleImageConfirmed"],
        "peopleImageUncertain": result["qa"]["peopleImageUncertain"],
        "peopleImageMissing": result["qa"]["peopleImageMissing"],
        "peopleImageRejected": result["qa"]["peopleImageRejected"],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
