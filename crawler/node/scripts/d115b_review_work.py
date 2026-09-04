"""Generate D11.5B review suggestions from the candidate workbook.

This tool never edits the workbook, MySQL, Vue sources, or output/reviewed.
Its recommendations are advisory; every record remains unverified until a
human supplies an explicit decision in a later phase.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


SHEET_CONFIG = {
    "P0-历史人物": {"module": "people", "prefix": "PEOPLE", "file": "people-review.json", "name": "姓名"},
    "P0-历史事件": {"module": "history", "prefix": "HISTORY", "file": "history-events-review.json", "name": "事件名称"},
    "P0-七贤庄院落": {"module": "courtyard", "prefix": "COURTYARD", "file": "courtyards-review.json", "name": "院落"},
    "P1-参观服务": {"module": "visit", "prefix": "VISIT", "file": "visit-review.json", "name": "项目"},
    "P1-文物": {"module": "relic", "prefix": "RELIC", "file": "relics-review.json", "name": "文物名称"},
    "P2-展览": {"module": "exhibition", "prefix": "EXHIBITION", "file": "exhibitions-review.json", "name": "展览名称"},
    "P2-新闻": {"module": "news", "prefix": "NEWS", "file": "news-review.json", "name": "标题"},
    "P2-馆情": {"module": "museum", "prefix": "MUSEUM", "file": "museum-review.json", "name": "项目"},
}

MODULE_LABELS = {
    "people": "人物", "history": "历史", "courtyard": "院落", "visit": "参观服务",
    "relic": "文物", "exhibition": "展览", "news": "新闻", "museum": "馆情",
}

EXPECTED_COUNTS = {
    "people": 31, "history": 23, "courtyard": 7, "visit": 10,
    "relic": 17, "exhibition": 11, "news": 12, "museum": 11,
}


def candidate_id_for(sheet_name: str, sequence: int) -> str:
    """Return a workflow-only ID that keeps the workbook sequence stable."""
    return f"{SHEET_CONFIG[sheet_name]['prefix']}-{int(sequence):03d}"


def parse_time_fields(value: object) -> dict:
    """Keep literal time text and only parse explicit year/month/day fields."""
    text = "" if value is None else str(value).strip()
    result = {"timeText": text, "year": None, "month": None, "day": None, "precision": "unknown"}
    day_match = re.search(r"(?P<year>\d{4})年(?P<month>\d{1,2})月(?P<day>\d{1,2})日", text)
    if day_match:
        result.update({key: int(day_match.group(key)) for key in ("year", "month", "day")})
        result["precision"] = "day"
        return result
    month_match = re.search(r"(?P<year>\d{4})年(?P<month>\d{1,2})月", text)
    if month_match:
        result["year"] = int(month_match.group("year"))
        result["month"] = int(month_match.group("month"))
        result["precision"] = "month"
        return result
    year_match = re.search(r"(?P<year>\d{4})年", text)
    if year_match:
        result["year"] = int(year_match.group("year"))
        result["precision"] = "year"
    return result


def source_level_for(source_type: str) -> str:
    text = source_type or ""
    if any(token in text for token in ("官网", "国家保密局", "地方志", "文物局", "革命文物教育平台")):
        return "A"
    if any(token in text for token in ("人民网", "党史频道", "央视网", "纪念网")):
        return "B"
    return "C"


def recommend_decision(module: str, source_level: str, source_url: str, original_status: str, has_core_fields: bool, weak_only: bool) -> tuple[str, str]:
    """Return an advisory decision, never a final human approval."""
    status = original_status or ""
    if any(token in status for token in ("口径差异", "两处时间", "撤离月日", "年份口径")):
        return "CONFLICT", "Excel 原始核验状态已标注口径差异，需人工裁决后才能采用。"
    if not source_url:
        return "PENDING", "缺少可追溯的来源 URL，不能形成可批准候选。"
    if module == "courtyard" and "用途待核实" in status:
        return "PENDING", "公告只提供院落线索，具体用途仍待人工核实。"
    if weak_only:
        if any(token in status for token in ("待核实", "暂缓", "参考信息")):
            return "REJECT", "仅有弱来源且原始状态已建议暂缓/参考，不建议进入正式候选。"
        return "PENDING", "当前只有弱来源支撑，需补充 A/B 级来源。"
    if not has_core_fields:
        return "PENDING", "核心名称、关系、时间或内容字段不完整，需补充后再审。"
    if module == "people" and "文物线索" in status:
        return "PENDING", "来源支持文物线索，但人物关系与详情尚需补充。"
    if module == "relic" and "补充线索" in status:
        return "PENDING", "补充线索尚未明确证明为馆藏正式记录，暂不建议批准。"
    if module == "visit" and "最新公告" in status:
        return "PENDING", "参观服务具时效性，需以人工确认时的最新官方公告为准。"
    if module == "visit" and "媒体报道" in status:
        return "PENDING", "讲解服务仅由媒体报道支撑，需补充当前官方说明。"
    if module == "people":
        return "APPROVE", "A/B 级来源支持人物文字关系；图片可保持为空，不影响文字资料的人工审批。"
    return "APPROVE", "来源和核心字段满足审核建议条件；仍需由人工作出最终决定。"


def normalized(value: object) -> str:
    return re.sub(r"[\s《》“”'\"（）()、，。·—_-]", "", "" if value is None else str(value)).lower()


def read_rows(workbook_path: Path) -> dict[str, list[dict]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=True)
    by_sheet = {}
    for sheet_name in SHEET_CONFIG:
        worksheet = workbook[sheet_name]
        headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
        rows = []
        for row_index in range(2, worksheet.max_row + 1):
            row = {headers[column - 1]: worksheet.cell(row_index, column).value for column in range(1, worksheet.max_column + 1)}
            if any(value is not None for value in row.values()):
                rows.append(row)
        by_sheet[sheet_name] = rows
    return by_sheet


def missing_fields(module: str, row: dict, image_status: str) -> list[str]:
    missing = []
    if not row.get("来源(URL)"):
        missing.append("sourceUrl")
    if module == "people":
        if not row.get("与西安八办/七贤庄的关系"):
            missing.append("relationship")
        if image_status != "confirmed":
            missing.append("reliablePortrait")
    elif module == "history":
        if not row.get("时间"):
            missing.append("timeText")
        if not row.get("事件描述/关键细节"):
            missing.append("description")
    elif module == "courtyard" and not row.get("名称/用途"):
        missing.append("historicalUse")
    elif module == "relic":
        if not row.get("类别（补全）"):
            missing.append("category")
        if not row.get("年代（补全）"):
            missing.append("era")
        if image_status != "confirmed":
            missing.append("confirmedImage")
    elif module == "news":
        missing.extend(["detailSourceUrl", "confirmedCoverImage"])
    return missing


def image_status_for(module: str, row: dict) -> str:
    if module == "people":
        text = str(row.get("图片来源") or "")
        return "missing" if not text or "待补" in text else "uncertain"
    if module == "relic":
        text = str(row.get("图片关联") or "")
        return "uncertain" if text else "missing"
    return "missing"


def exhibition_type_for(value: object) -> str:
    text = str(value or "")
    if "基本陈列" in text:
        return "basicCollection"
    if "旧址复原" in text:
        return "siteRestoration"
    if "数字" in text or "元宇宙" in text:
        return "digitalExhibition"
    if "专题" in text:
        return "specialExhibition"
    return "unknown"


def visit_rule_type_for(row: dict) -> str:
    text = f"{row.get('项目') or ''} {row.get('内容') or ''}"
    if any(token in text for token in ("五一", "清明", "春节", "节假日", "延时")):
        return "HOLIDAY_NOTICE"
    if "开放时间" in text or "周一" in text:
        return "REGULAR"
    if "公告" in text:
        return "TEMPORARY_NOTICE"
    return "NOT_APPLICABLE"


def build_review_record(sheet_name: str, row: dict) -> dict:
    config = SHEET_CONFIG[sheet_name]
    module = config["module"]
    sequence = int(row["序号"])
    source_type = str(row.get("来源类型") or "")
    source_url = str(row.get("来源(URL)") or "")
    original_status_key = "时效性/核验状态" if module == "visit" else "核验状态"
    original_status = str(row.get(original_status_key) or "")
    source_level = source_level_for(source_type)
    weak_only = source_level == "C"
    name = str(row.get(config["name"]) or "")
    image_status = image_status_for(module, row)
    core_fields = bool(name and source_url)
    if module == "people":
        core_fields = core_fields and bool(row.get("与西安八办/七贤庄的关系"))
    if module == "history":
        core_fields = core_fields and bool(row.get("时间") and row.get("事件描述/关键细节"))
    if module == "courtyard":
        core_fields = core_fields and bool(row.get("名称/用途"))
    decision, reason = recommend_decision(module, source_level, source_url, original_status, core_fields, weak_only)
    conflict = decision == "CONFLICT"
    record = {
        "candidateId": candidate_id_for(sheet_name, sequence),
        "module": module,
        "originalRow": row,
        "originalText": {
            "name": name,
            "sourceType": source_type,
            "sourceUrl": source_url,
            "sourceEvidenceStatus": original_status,
        },
        "sourceEvidenceStatus": "strong" if source_level in {"A", "B"} else "weak",
        "sourceLevel": source_level,
        "sourceAvailability": "notChecked",
        "conflictStatus": "conflict" if conflict else "none",
        "conflictingValues": [row.get("时间"), row.get("事件描述/关键细节")] if conflict else [],
        "conflictingSources": [source_url] if conflict and source_url else [],
        "missingFields": missing_fields(module, row, image_status),
        "imageStatus": image_status,
        "recommendedDecision": decision,
        "recommendationReason": reason,
        "reviewDecision": None,
        "verified": False,
        "reviewStatus": "pending",
    }
    if module == "people":
        record.update({
            "originalRelation": row.get("与西安八办/七贤庄的关系"),
            "originalDetails": row.get("关键细节"),
            "relationshipVerified": source_level in {"A", "B"} and bool(row.get("与西安八办/七贤庄的关系")),
            "roleVerified": source_level in {"A", "B"} and bool(row.get("身份/职务")),
            "detailsVerified": source_level in {"A", "B"} and bool(row.get("关键细节")),
            "imageVerified": False,
        })
    elif module == "history":
        record.update(parse_time_fields(row.get("时间")))
        record["originalDetails"] = row.get("事件描述/关键细节")
    elif module == "courtyard":
        record["entityType"] = "siteOverview" if "整体" in name else "courtyard"
        record["originalUse"] = row.get("名称/用途")
    elif module == "visit":
        record["visitRuleType"] = visit_rule_type_for(row)
    elif module == "relic":
        record["grade"] = "国家一级革命文物" if "一级革命文物" in f"{row.get('年代（补全）') or ''} {original_status}" else None
        record["gradeSource"] = source_url if record["grade"] else None
    elif module == "exhibition":
        record["exhibitionType"] = exhibition_type_for(row.get("类型"))
        record["exhibitionStatus"] = "unknown"
    return record


def detect_duplicates(records: list[dict]) -> list[dict]:
    grouped = defaultdict(list)
    for record in records:
        row = record["originalRow"]
        module = record["module"]
        config = next(value for value in SHEET_CONFIG.values() if value["module"] == module)
        if module == "history":
            key = f"{normalized(row.get('时间'))}|{normalized(row.get('事件名称'))}|{normalized(row.get('事件描述/关键细节'))}"
        elif module == "courtyard":
            key = f"{normalized(row.get('院落'))}|{normalized(row.get('名称/用途'))}"
        elif module in {"relic", "exhibition", "news"}:
            key = f"{normalized(row.get(config['name']))}|{normalized(row.get('来源(URL)'))}"
        else:
            key = normalized(row.get(config["name"]))
        grouped[(module, key)].append(record["candidateId"])
    return [
        {"module": module, "candidateIds": ids}
        for (module, _), ids in grouped.items()
        if len(ids) > 1
    ]


def batch_for(record: dict) -> str:
    batch_one = {
        "PEOPLE-001", "PEOPLE-002", "PEOPLE-003", "PEOPLE-004", "PEOPLE-005",
        "HISTORY-003", "HISTORY-005", "HISTORY-014", "HISTORY-015", "HISTORY-022",
        "COURTYARD-002", "COURTYARD-004", "COURTYARD-005", "COURTYARD-006",
        "MUSEUM-001", "MUSEUM-002", "MUSEUM-003", "MUSEUM-004",
        *{f"RELIC-{number:03d}" for number in range(1, 10)}, "EXHIBITION-001",
    }
    if record["candidateId"] in batch_one:
        return "Batch 1"
    if record["recommendedDecision"] in {"CONFLICT", "REJECT"} or record["sourceEvidenceStatus"] == "weak":
        return "Batch 3"
    return "Batch 2"


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_summary(docs_dir: Path, records: list[dict], duplicates: list[dict]) -> None:
    decision_counts = Counter(record["recommendedDecision"] for record in records)
    module_counts = Counter(record["module"] for record in records)
    batch_counts = Counter(record["batch"] for record in records)
    weak_only = sum(record["sourceEvidenceStatus"] == "weak" for record in records)
    missing_images = sum(record["module"] == "people" and record["imageStatus"] != "confirmed" for record in records)
    lines = [
        "# D11.5B 候选资料审核建议汇总", "",
        "本报告的 APPROVE/PENDING/CONFLICT/REJECT 全部是 Codex 的**建议**，不是人工最终决定。所有记录保持 `verified: false`、`reviewDecision: null`、`reviewStatus: pending`；本轮未生成或修改正式 `reviewed/` 数据。", "",
        "## 总量与模块", "", "| 模块 | Excel 实际记录数 | 审核工作记录数 |", "| --- | ---: | ---: |",
    ]
    for module in EXPECTED_COUNTS:
        lines.append(f"| {MODULE_LABELS[module]} | {EXPECTED_COUNTS[module]} | {module_counts[module]} |")
    lines.extend(["", f"总候选：{len(records)}。", "", "## 建议决策统计", "", "| 建议 | 数量 |", "| --- | ---: |"])
    for decision in ("APPROVE", "PENDING", "CONFLICT", "REJECT"):
        lines.append(f"| {decision} | {decision_counts[decision]} |")
    lines.extend([
        "", "## 质量检查", "",
        f"- 潜在重复：{len(duplicates)}。", f"- 来源冲突：{decision_counts['CONFLICT']}。", f"- 仅弱来源支撑：{weak_only}。",
        f"- 人物缺少可靠确认肖像：{missing_images}。", "- 最常见缺口：人物可靠肖像、新闻详情页稳定 URL/确认封面、补充文物的馆藏身份与图片、参观服务的当前有效性。",
        "", "## 批次", "", f"- Batch 1：{batch_counts['Batch 1']} 条，优先审核关系直接人物、无冲突关键节点、1/3/4/7 号院、核心馆情、原 9 件文物和基本陈列。",
        f"- Batch 2：{batch_counts['Batch 2']} 条，一般候选与需要补字段的官方资料。",
        f"- Batch 3：{batch_counts['Batch 3']} 条，冲突、弱来源或明确暂缓/拒绝建议。",
        "", "## 处理边界", "", "- Excel 原始文字保留于每条 `originalRow` / `originalText`，未润色为网站正文。",
        "- 历史时间仅解析明示年、月、日；如“1934年冬”保持原文且只标记 year precision。",
        "- 人物图片未因文字来源可靠而自动确认；只有人工确认身份和原始图片页后才可标记 confirmed。",
        "- 参观服务的常规规则、节假日公告、临时公告在 review-work 中分开标识，预约方式差异保持 conflict。",
    ])
    (docs_dir / "d11-5b-review-summary.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_manual_list(docs_dir: Path, records: list[dict]) -> None:
    lines = [
        "# D11.5B 人工决策清单", "",
        "请仅在最后一列填写：批准、暂缓、冲突或拒绝。填写前不应把任何 recommendedDecision 当作最终审核结论。", "",
    ]
    for batch in ("Batch 1", "Batch 2", "Batch 3"):
        lines.extend([f"## {batch}", "", "| ID | 模块 | 名称 | 推荐 | 主要证据 | 缺失/冲突 | 我的决策 |", "| --- | --- | --- | --- | --- | --- | --- |"])
        for record in (item for item in records if item["batch"] == batch):
            row = record["originalRow"]
            config = next(value for value in SHEET_CONFIG.values() if value["module"] == record["module"])
            name = str(row.get(config["name"]) or "").replace("|", "／")
            evidence = f"{record['sourceLevel']}级；{row.get('来源类型') or '未说明'}".replace("|", "／")
            gaps = "、".join(record["missingFields"]) or ("来源冲突" if record["conflictStatus"] == "conflict" else "无")
            lines.append(f"| {record['candidateId']} | {MODULE_LABELS[record['module']]} | {name} | {record['recommendedDecision']} | {evidence} | {gaps} |  |")
        lines.append("")
    (docs_dir / "manual-approval-list.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def generate(workbook_path: Path, output_dir: Path, docs_dir: Path) -> dict:
    rows_by_sheet = read_rows(workbook_path)
    for sheet_name, expected in EXPECTED_COUNTS.items():
        matching_sheet = next(name for name, config in SHEET_CONFIG.items() if config["module"] == sheet_name)
        if len(rows_by_sheet[matching_sheet]) != expected:
            raise ValueError(f"{matching_sheet} expected {expected} records, got {len(rows_by_sheet[matching_sheet])}")
    output_dir.mkdir(parents=True, exist_ok=True)
    docs_dir.mkdir(parents=True, exist_ok=True)
    records = []
    for sheet_name, config in SHEET_CONFIG.items():
        module_records = [build_review_record(sheet_name, row) for row in rows_by_sheet[sheet_name]]
        records.extend(module_records)
        write_json(output_dir / config["file"], module_records)
    duplicates = detect_duplicates(records)
    for record in records:
        record["duplicateCandidates"] = [group["candidateIds"] for group in duplicates if record["candidateId"] in group["candidateIds"]]
        record["batch"] = batch_for(record)
    for config in SHEET_CONFIG.values():
        module_records = [record for record in records if record["module"] == config["module"]]
        write_json(output_dir / config["file"], module_records)
    qa = {
        "total": len(records),
        "missingSourceUrl": sum(not record["originalText"]["sourceUrl"] for record in records),
        "potentialDuplicates": len(duplicates),
        "duplicateCandidates": duplicates,
        "conflicts": sum(record["conflictStatus"] == "conflict" for record in records),
        "weakSourceOnly": sum(record["sourceEvidenceStatus"] == "weak" for record in records),
        "missingImages": sum(record["imageStatus"] == "missing" for record in records),
        "uncertainImages": sum(record["imageStatus"] == "uncertain" for record in records),
        "confirmedImages": sum(record["imageStatus"] == "confirmed" for record in records),
    }
    write_json(output_dir / "qa-summary.json", qa)
    write_summary(docs_dir, records, duplicates)
    write_manual_list(docs_dir, records)
    return {"records": records, "qa": qa}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--docs-dir", required=True, type=Path)
    args = parser.parse_args()
    result = generate(args.input, args.output_dir, args.docs_dir)
    print(json.dumps({"total": len(result["records"]), **result["qa"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
